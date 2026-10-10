# importar las librerias para leer el excel y mostrar la tabla
import openpyxl
from tabulate import tabulate

# abrir el archivo de excel y tomamos la hoja activa
excel_dataframe = openpyxl.load_workbook("personas.xlsx")
dataframe = excel_dataframe.active

# lista donde se guarda toda la informacion del excel
data = []

# recorrer las filas del excel
for row in range(1, dataframe.max_row):
    _row = [row]

    # recorremos las columnas de la fila y guardamos cada valor
    for col in dataframe.iter_cols(1, dataframe.max_column):
        _row.append(col[row].value)

    # guarda la fila completa en la lista principal
    data.append(_row)

# encabezados de la tabla
headers = ["#", "Id", "Name", "Company", "Email", "MAC Address"]
headers_align = ("center",) * 6


# funcion para mostrar las tablas
def mostrar(lista):
    print(tabulate(lista, headers=headers, tablefmt="fancy_grid", colalign=headers_align))


# tabla completa
mostrar(data)

# reto 1: solo Name y Email
datos_reto1 = []
for p in data:
    datos_reto1.append([p[2], p[4]])
print(tabulate(datos_reto1, headers=["Name", "Email"], tablefmt="fancy_grid"))

# reto 2: contar registros
print("Total de registros:", len(data))


# reto 5: menu interactivo (incluye los retos 3 y 4)
opcion = ""
while opcion != "5":
    print("")
    print("1. Ver todos")
    print("2. Buscar por ID")
    print("3. Filtrar empresa")
    print("4. Contar registros")
    print("5. Salir")
    opcion = input("Elige una opcion: ")

    if opcion == "1":
        mostrar(data)

    elif opcion == "2":
        # reto 4: buscar una persona por su ID
        id_buscado = input("Escribe el ID: ")
        # validamos que el ID sea un numero
        if id_buscado.isdigit() == False:
            print("El ID debe ser un numero")
        else:
            encontrados = []
            for p in data:
                if str(p[1]) == id_buscado:
                    encontrados.append(p)
            if len(encontrados) == 0:
                print("No se encontro ese ID")
            else:
                mostrar(encontrados)

    elif opcion == "3":
        # reto 3: filtrar las personas por la empresa
        palabra = input("Escribe la palabra de la empresa (ejemplo Group): ")
        filtrados = []
        for p in data:
            # pasamos todo a minuscula para que no importe como se escriba
            if palabra.lower() in str(p[3]).lower():
                filtrados.append(p)
        if len(filtrados) == 0:
            print("No hay empresas con esa palabra")
        else:
            mostrar(filtrados)

    elif opcion == "4":
        print("Total de registros:", len(data))

    elif opcion == "5":
        print("Programa finalizado")

    else:
        print("Opcion no valida")